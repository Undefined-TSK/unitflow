"""离线文档解析（成熟库 + 安全防护）。

遵循「零自造轮子」：
- 网页/HTML 正文：trafilatura（F-Score 0.926 SOTA）主 + BeautifulSoup 兜底；
- PDF：pypdfium2（快）主 + pdfminer.six 兜底；
- DOCX：mammoth（正文干净）主 + python-docx 兜底；
- PPTX：python-pptx；XLSX：openpyxl；EPUB：ebooklib。
此模块只做「单文件 → 纯文本」这一职责，不含抓取与入库逻辑。
"""
from __future__ import annotations

import zipfile
from pathlib import Path

from ..config import settings


# 离线文档支持格式
SUPPORTED_EXTS = [".pdf", ".epub", ".md", ".markdown", ".txt", ".rst",
                  ".html", ".htm", ".docx", ".pptx", ".xlsx"]


def html_to_text(html: str | bytes) -> str:
    """用 trafilatura 抽取网页/HTML 正文（含表格），去掉导航/脚本等噪声。"""
    try:
        import trafilatura

        text = trafilatura.extract(
            html, include_comments=False, include_tables=True,
            include_formatting=False, favor_recall=True,
        )
        if text and text.strip():
            return text
    except Exception:
        pass
    # 兜底：BeautifulSoup 纯文本
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html if isinstance(html, str) else "", "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    return soup.get_text("\n")


def _check_zip_safe(path: str, max_ratio: float = 100.0) -> None:
    """ZIP 炸弹防护：EPUB/DOCX/XLSX/PPTX 均为 zip，压缩比 >100:1 拒绝。"""
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if info.file_size and info.compress_size:
                ratio = info.file_size / max(info.compress_size, 1)
                if ratio > max_ratio:
                    raise ValueError(f"压缩比异常 {ratio:.0f}:1，疑似解压炸弹，已拒绝解析")


def _pdf_text(path: str) -> str:
    """PDF 抽取：pypdfium2（快）主 + pdfminer.six（全）兜底。"""
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(path)
        parts = []
        for page in pdf:
            tp = page.get_textpage()
            parts.append(tp.get_text_range() or "")
        text = "\n".join(parts)
        if text.strip():
            return text
    except Exception:
        pass
    from pdfminer.high_level import extract_text

    return extract_text(path)


def _docx_text(path: str) -> str:
    """DOCX：mammoth（正文更干净）主 + python-docx 兜底。"""
    try:
        import mammoth

        with open(path, "rb") as f:
            result = mammoth.convert_to_markdown(f)
        text = result.value or ""
        if text.strip():
            return text
    except Exception:
        pass
    import docx

    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def _pptx_text(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# 第 {i} 页")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
    return "\n\n".join(parts)


def _xlsx_text(path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# 工作表 {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n\n".join(parts)


def _epub_text(path: str) -> str:
    import ebooklib
    from ebooklib import epub

    book = epub.read_epub(path)
    out = []
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        out.append(html_to_text(item.get_content()))
    return "\n\n".join(out)


def parse_file(path: str, filename: str | None = None) -> str:
    """离线文档解析，全部使用成熟库。调用方应先用 check_file_safe 校验。"""
    ext = Path(filename or path).suffix.lower()
    try:
        if ext in (".md", ".markdown", ".txt", ".rst"):
            return Path(path).read_text(encoding="utf-8", errors="ignore")
        if ext in (".html", ".htm"):
            return html_to_text(Path(path).read_text(encoding="utf-8", errors="ignore"))
        if ext == ".pdf":
            return _pdf_text(path)
        if ext == ".docx":
            return _docx_text(path)
        if ext == ".pptx":
            return _pptx_text(path)
        if ext == ".xlsx":
            return _xlsx_text(path)
        if ext == ".epub":
            return _epub_text(path)
    except Exception as e:  # noqa: BLE001
        return f"[解析失败] {e}"
    return "[不支持的格式]"


def check_file_safe(path: str) -> str:
    """桌面安全防护：大小上限 + ZIP 炸弹。返回空字符串表示安全，否则为拒绝原因。"""
    try:
        size_mb = Path(path).stat().st_size / 1024 / 1024
    except OSError as e:
        return f"无法读取文件：{e}"
    if size_mb > settings.MAX_UPLOAD_MB:
        return f"文件 {size_mb:.1f}MB 超过上限 {settings.MAX_UPLOAD_MB:.0f}MB"
    if Path(path).suffix.lower() in (".epub", ".docx", ".pptx", ".xlsx"):
        try:
            _check_zip_safe(path)
        except ValueError as e:
            return str(e)
        except Exception:
            return "压缩包损坏，无法解析"
    return ""
